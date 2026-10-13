import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen slides (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette (Ronas IT / Learnify Theme)
    CANVAS_BG = RGBColor(213, 223, 232)      # #D5DFE8
    WHITE = RGBColor(255, 255, 255)          # #FFFFFF
    CHARCOAL = RGBColor(24, 25, 30)          # #18191E
    YELLOW = RGBColor(254, 232, 131)         # #FEE883 (Sunny Yellow)
    LAVENDER = RGBColor(224, 212, 253)       # #E0D4FD (Pastel Lavender)
    SKY = RGBColor(186, 230, 253)            # #BAE6FD (Pastel Sky)
    CORAL = RGBColor(255, 91, 55)            # #FF5B37 (Coral Accent)
    AMBER = RGBColor(255, 210, 51)           # #FFD233
    MUTED = RGBColor(100, 116, 139)          # #64748B
    CARD_BORDER = RGBColor(226, 232, 240)

    blank_layout = prs.slide_layouts[6]

    def set_canvas_background(slide):
        # Add canvas rectangle covering entire slide
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = CANVAS_BG
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, fill_color, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_badge(slide, left, top, text, bg_color=YELLOW, text_color=CHARCOAL):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.2), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_color
        badge.line.fill.background()
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = text_color
        p.font.name = "Arial"
        return badge

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide1)

    # Master Tablet Frame Card
    add_card(slide1, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)

    # Badge
    add_badge(slide1, Inches(1.3), Inches(1.1), "✨ AI CAREER NAVIGATOR", YELLOW, CHARCOAL)

    # Title & Coral Accent
    tx_box = slide1.shapes.add_textbox(Inches(1.3), Inches(1.7), Inches(10.5), Inches(2.2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "PathPilot AI"
    p1.font.size = Pt(54)
    p1.font.bold = True
    p1.font.color.rgb = CHARCOAL
    p1.font.name = "Arial"

    p2 = tf.add_paragraph()
    p2.text = "Your Personalized AI Career & Skill Navigator"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = CORAL
    p2.font.name = "Arial"

    p3 = tf.add_paragraph()
    p3.text = "Transforming career confusion into an inspiring, step-by-step learning journey with zero guesswork."
    p3.font.size = Pt(14)
    p3.font.color.rgb = MUTED
    p3.font.name = "Arial"

    # Founders Bento Strip
    f_card = add_card(slide1, Inches(1.3), Inches(4.3), Inches(10.7), Inches(1.8), CHARCOAL)
    f_box = slide1.shapes.add_textbox(Inches(1.5), Inches(4.45), Inches(10.3), Inches(1.5))
    ftf = f_box.text_frame
    ftf.word_wrap = True
    fp1 = ftf.paragraphs[0]
    fp1.text = "FOUNDED & ENGINEERED BY TEAM PATHPILOT"
    fp1.font.size = Pt(11)
    fp1.font.bold = True
    fp1.font.color.rgb = AMBER

    fp2 = ftf.add_paragraph()
    fp2.text = "👩‍💻 Sobia Sattar (AI Systems Lead)   •   👨‍💻 Muhammad Hassan (Data Systems Lead)   •   👩‍🎨 Sumra Shar (Creative Product Lead)"
    fp2.font.size = Pt(14)
    fp2.font.bold = True
    fp2.font.color.rgb = WHITE

    fp3 = ftf.add_paragraph()
    fp3.text = "Featuring Pip the Cute Multilingual Buddy  |  Built with Python 3.14, FastAPI, React & Vite"
    fp3.font.size = Pt(11)
    fp3.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: THE PROBLEM (WHY WE BUILT IT)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide2)
    add_card(slide2, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide2, Inches(1.3), Inches(0.95), "THE CHALLENGE", LAVENDER, CHARCOAL)

    h_box = slide2.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    htf = h_box.text_frame
    hp = htf.paragraphs[0]
    hp.text = "Why Do Tech Careers Feel So Overwhelming?"
    hp.font.size = Pt(30)
    hp.font.bold = True
    hp.font.color.rgb = CHARCOAL

    # 3 Pastel Bento Problem Cards
    probs = [
        ("01", "Confusing Job Requirements", "Entry-level jobs asking for 5+ years of experience and 20 different technologies. Students don't know what's actually required vs optional.", YELLOW),
        ("02", "No Personalized Direction", "Online courses and advice are generic. Students waste hundreds of hours studying random tutorials without knowing their actual competency gaps.", LAVENDER),
        ("03", "Robotic & Cold Guidance", "Existing tools produce intimidating, clinical compliance reports that discourage learners instead of providing positive, actionable steps.", SKY)
    ]
    card_w = Inches(3.45)
    for idx, (num, title, desc, col) in enumerate(probs):
        c_left = Inches(1.3) + idx * Inches(3.62)
        add_card(slide2, c_left, Inches(2.4), card_w, Inches(3.8), col)
        
        c_box = slide2.shapes.add_textbox(c_left + Inches(0.2), Inches(2.6), card_w - Inches(0.4), Inches(3.3))
        ctf = c_box.text_frame
        ctf.word_wrap = True
        
        p_num = ctf.paragraphs[0]
        p_num.text = f"PROBLEM {num}"
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = CHARCOAL
        
        p_t = ctf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = CHARCOAL
        
        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 3: OUR SOLUTION (PATHPILOT OVERVIEW)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide3)
    add_card(slide3, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide3, Inches(1.3), Inches(0.95), "OUR SOLUTION", SKY, CHARCOAL)

    h_box3 = slide3.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp3 = h_box3.text_frame.paragraphs[0]
    hp3.text = "PathPilot AI: A Friendly, Honest Career Companion"
    hp3.font.size = Pt(30)
    hp3.font.bold = True
    hp3.font.color.rgb = CHARCOAL

    sol_cards = [
        ("🎯", "Accurate Skill Scanner", "Extracts real skills from your resume with PyMuPDF. Zero fabricated data with 100% student confirmation control.", YELLOW),
        ("📊", "Interactive Growth Map", "Visualizes your level (0 to 4) against industry standards. 1-click level adjustments make updating effortless.", LAVENDER),
        ("🗺️", "Dynamic 4-Phase Roadmap", "Video lectures, milestone projects, and live code sandbox practice calibrated to your weekly study hours.", SKY),
        ("🦊", "Pip: Multilingual AI Buddy", "A cute mascot ready 24/7 to answer questions and explain concepts in English, Urdu (اردو), or Roman Urdu.", YELLOW)
    ]
    sw = Inches(5.2)
    sh = Inches(1.8)
    for idx, (icon, title, desc, col) in enumerate(sol_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        s_left = Inches(1.3) + col_idx * Inches(5.5)
        s_top = Inches(2.4) + row_idx * Inches(2.0)
        
        add_card(slide3, s_left, s_top, sw, sh, col)
        
        box = slide3.shapes.add_textbox(s_left + Inches(0.2), s_top + Inches(0.15), sw - Inches(0.4), sh - Inches(0.3))
        btf = box.text_frame
        btf.word_wrap = True
        
        bp1 = btf.paragraphs[0]
        bp1.text = f"{icon}  {title}"
        bp1.font.size = Pt(16)
        bp1.font.bold = True
        bp1.font.color.rgb = CHARCOAL
        
        bp2 = btf.add_paragraph()
        bp2.text = desc
        bp2.font.size = Pt(11.5)
        bp2.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 4: MEET THE FOUNDERS
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide4)
    add_card(slide4, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide4, Inches(1.3), Inches(0.95), "TEAM PATHPILOT", YELLOW, CHARCOAL)

    h_box4 = slide4.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp4 = h_box4.text_frame.paragraphs[0]
    hp4.text = "Meet the Creators: Crafted with Passion & Purpose"
    hp4.font.size = Pt(30)
    hp4.font.bold = True
    hp4.font.color.rgb = CHARCOAL

    founders = [
        ("👩‍💻 Sobia Sattar", "Co-Founder & AI Systems Lead", "Strategic Vision & AI Empathy",
         "Passionate about demystifying career growth. Sobia envisioned PathPilot as an intelligent, empathetic compass that translates confusing market requirements into confident next steps.",
         YELLOW),
        ("👨‍💻 Muhammad Hassan", "Co-Founder & Data Intelligence Lead", "Deep Systems Logic & Analytical Truth",
         "An analytical problem solver driven to eliminate guesswork. Hassan built the data engines and deterministic roadmap generators that ensure every metric is explainable and rock-solid.",
         LAVENDER),
        ("👩‍🎨 Sumra Shar", "Co-Founder & Creative Product Lead", "Creative Intuition & Delightful UX",
         "Believes learning tools should spark joy. Sumra crafted the signature Ronas IT pastel design system, Pip the cute mascot, and fluid animations that make planning feel playful.",
         SKY)
    ]
    fw = Inches(3.45)
    for idx, (name, role, power, bio, col) in enumerate(founders):
        f_left = Inches(1.3) + idx * Inches(3.62)
        add_card(slide4, f_left, Inches(2.4), fw, Inches(3.9), col)
        
        box = slide4.shapes.add_textbox(f_left + Inches(0.2), Inches(2.55), fw - Inches(0.4), Inches(3.6))
        btf = box.text_frame
        btf.word_wrap = True
        
        p1 = btf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = CHARCOAL
        
        p2 = btf.add_paragraph()
        p2.text = role
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = CHARCOAL
        
        p3 = btf.add_paragraph()
        p3.text = f"✨ Superpower: {power}"
        p3.font.size = Pt(10.5)
        p3.font.bold = True
        p3.font.color.rgb = CORAL
        
        p4 = btf.add_paragraph()
        p4.text = f"\n\"{bio}\""
        p4.font.size = Pt(11)
        p4.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 5: FEATURE 1 - RESUME & EXPERIENCE SCANNER
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide5)
    add_card(slide5, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide5, Inches(1.3), Inches(0.95), "FEATURE SPOTLIGHT 01", LAVENDER, CHARCOAL)

    h_box5 = slide5.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp5 = h_box5.text_frame.paragraphs[0]
    hp5.text = "Resume & Experience Scanner"
    hp5.font.size = Pt(30)
    hp5.font.bold = True
    hp5.font.color.rgb = CHARCOAL

    # Left Card: How it Works
    add_card(slide5, Inches(1.3), Inches(2.3), Inches(5.1), Inches(4.0), YELLOW)
    lbox = slide5.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(4.7), Inches(3.5))
    ltf = lbox.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "How the Scanner Works:"
    lp1.font.size = Pt(18)
    lp1.font.bold = True
    lp1.font.color.rgb = CHARCOAL

    bullets_l = [
        "📄 Upload any PDF resume or paste text directly.",
        "🔍 PyMuPDF extracts skills, education, and exact evidence sentences.",
        "⚡ Semantic alias mapping recognizes equivalent skills (e.g., 'React.js' = 'React').",
        "🔒 100% Private: Processed in isolated user sessions with zero external sharing."
    ]
    for b in bullets_l:
        p = ltf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = CHARCOAL

    # Right Card: The Human Review Layer
    add_card(slide5, Inches(6.7), Inches(2.3), Inches(5.3), Inches(4.0), CHARCOAL)
    rbox = slide5.shapes.add_textbox(Inches(6.9), Inches(2.5), Inches(4.9), Inches(3.5))
    rtf = rbox.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = "100% Student Confirmation Layer"
    rp1.font.size = Pt(18)
    rp1.font.bold = True
    rp1.font.color.rgb = WHITE

    bullets_r = [
        "✅ No AI Hallucinations: We never invent fake skills you didn't list.",
        "✏️ Easy Adjustments: Add missing skills or adjust proficiency with 1 tap.",
        "📊 Transparent Evidence: Hover over any skill to see the exact sentence from your resume.",
        "🎯 Confidence: Your roadmap is built on verified reality, not guesswork."
    ]
    for b in bullets_r:
        p = rtf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(226, 232, 240)

    # =========================================================================
    # SLIDE 6: FEATURE 2 - SKILL COMPASS & GROWTH TARGETS
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide6)
    add_card(slide6, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide6, Inches(1.3), Inches(0.95), "FEATURE SPOTLIGHT 02", SKY, CHARCOAL)

    h_box6 = slide6.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp6 = h_box6.text_frame.paragraphs[0]
    hp6.text = "Interactive Skill Compass & Growth Map"
    hp6.font.size = Pt(30)
    hp6.font.bold = True
    hp6.font.color.rgb = CHARCOAL

    f2_items = [
        ("🎨 Compact Bento Cards", "Replaced overwhelming, clinical audit lists with beautiful rounded cards displaying category tags and clean priorities.", YELLOW),
        ("📏 Unified 4-Segment Gauge", "Clear visual stepper showing your current level vs target level (e.g., Level 0 ➔ Level 3 needed).", LAVENDER),
        ("⚡ 1-Click Level Setter", "Click [0] [1] [2] [3] [4] directly on any skill card to self-calibrate your skills and recalculate readiness instantly.", SKY),
        ("💡 Positive Growth Tips", "Encouraging, actionable recommendations instead of alarming error warnings or robotic audits.", YELLOW)
    ]
    for idx, (title, desc, col) in enumerate(f2_items):
        col_idx = idx % 2
        row_idx = idx // 2
        s_left = Inches(1.3) + col_idx * Inches(5.5)
        s_top = Inches(2.4) + row_idx * Inches(2.0)
        
        add_card(slide6, s_left, s_top, Inches(5.2), Inches(1.8), col)
        
        box = slide6.shapes.add_textbox(s_left + Inches(0.2), s_top + Inches(0.15), Inches(4.8), Inches(1.5))
        btf = box.text_frame
        btf.word_wrap = True
        
        bp1 = btf.paragraphs[0]
        bp1.text = title
        bp1.font.size = Pt(16)
        bp1.font.bold = True
        bp1.font.color.rgb = CHARCOAL
        
        bp2 = btf.add_paragraph()
        bp2.text = desc
        bp2.font.size = Pt(12)
        bp2.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 7: FEATURE 3 - DYNAMIC LEARNING ROADMAP
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide7)
    add_card(slide7, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide7, Inches(1.3), Inches(0.95), "FEATURE SPOTLIGHT 03", YELLOW, CHARCOAL)

    h_box7 = slide7.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp7 = h_box7.text_frame.paragraphs[0]
    hp7.text = "Personalized 4-Phase Learning Roadmap"
    hp7.font.size = Pt(30)
    hp7.font.bold = True
    hp7.font.color.rgb = CHARCOAL

    # 4 Phases Horizontal Strip
    phases = [
        ("Phase 01", "Core Foundations", "Master essential languages & syntax", YELLOW),
        ("Phase 02", "Applied Architecture", "Build real data & backend pipelines", LAVENDER),
        ("Phase 03", "Advanced Optimization", "Fine-tuning, scaling & testing", SKY),
        ("Phase 04", "Portfolio Deliverable", "Deploy production project for hiring", YELLOW)
    ]
    pw = Inches(2.55)
    for idx, (p_num, p_title, p_sub, p_col) in enumerate(phases):
        p_left = Inches(1.3) + idx * Inches(2.72)
        add_card(slide7, p_left, Inches(2.4), pw, Inches(1.8), p_col)
        
        box = slide7.shapes.add_textbox(p_left + Inches(0.15), Inches(2.5), pw - Inches(0.3), Inches(1.6))
        btf = box.text_frame
        btf.word_wrap = True
        
        bp1 = btf.paragraphs[0]
        bp1.text = p_num
        bp1.font.size = Pt(12)
        bp1.font.bold = True
        bp1.font.color.rgb = CHARCOAL
        
        bp2 = btf.add_paragraph()
        bp2.text = p_title
        bp2.font.size = Pt(14)
        bp2.font.bold = True
        bp2.font.color.rgb = CHARCOAL
        
        bp3 = btf.add_paragraph()
        bp3.text = p_sub
        bp3.font.size = Pt(10)
        bp3.font.color.rgb = CHARCOAL

    # Bottom features
    add_card(slide7, Inches(1.3), Inches(4.5), Inches(10.7), Inches(1.8), CHARCOAL)
    bbox = slide7.shapes.add_textbox(Inches(1.5), Inches(4.65), Inches(10.3), Inches(1.5))
    bbtf = bbox.text_frame
    bbtf.word_wrap = True
    
    bbp1 = bbtf.paragraphs[0]
    bbp1.text = "BUILT-IN INTERACTIVE WORKSHOP & STUDY INTEGRITY GATES"
    bbp1.font.size = Pt(12)
    bbp1.font.bold = True
    bbp1.font.color.rgb = AMBER
    
    bbp2 = bbtf.add_paragraph()
    bbp2.text = "• Real Playable Video Player: High-quality curated lectures with chapter timestamps.\n• Live Code Sandbox: Run Python and PyTorch test suites directly in your browser.\n• Verification Gate: Prevents accidental completion skips until you actually watch or code!"
    bbp2.font.size = Pt(12)
    bbp2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 8: FEATURE 4 - MEET PIP: MULTILINGUAL AI BUDDY
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide8)
    add_card(slide8, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide8, Inches(1.3), Inches(0.95), "FEATURE SPOTLIGHT 04", CORAL, WHITE)

    h_box8 = slide8.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp8 = h_box8.text_frame.paragraphs[0]
    hp8.text = "Meet Pip: Your Multilingual AI Career Buddy 🦊"
    hp8.font.size = Pt(30)
    hp8.font.bold = True
    hp8.font.color.rgb = CHARCOAL

    pip_features = [
        ("🐾 Cute Mascot Design", "Matches the Ronas IT / Learnify theme with headphone avatar, rosy blush, and sunny-yellow branding. Pops up automatically to greet learners.", YELLOW),
        ("🗣️ True Multilingual NLP", "Ask questions in English, Urdu (اردو with full RTL script), or Roman Urdu ('Score kaise barhayen?'). Pip responds in the same language naturally.", LAVENDER),
        ("🧠 High-Precision Answers", "Accurately explains complex technical topics: Gradient Descent, Overfitting, Precision vs Recall, SQL Joins, React Virtual DOM, and scoring math.", SKY),
        ("💡 24/7 Encouragement", "Tailors advice directly to your target role, readiness score, and current milestone. Offers interview tips and daily motivation.", YELLOW)
    ]
    for idx, (title, desc, col) in enumerate(pip_features):
        col_idx = idx % 2
        row_idx = idx // 2
        s_left = Inches(1.3) + col_idx * Inches(5.5)
        s_top = Inches(2.4) + row_idx * Inches(2.0)
        
        add_card(slide8, s_left, s_top, Inches(5.2), Inches(1.8), col)
        
        box = slide8.shapes.add_textbox(s_left + Inches(0.2), s_top + Inches(0.15), Inches(4.8), Inches(1.5))
        btf = box.text_frame
        btf.word_wrap = True
        
        bp1 = btf.paragraphs[0]
        bp1.text = title
        bp1.font.size = Pt(16)
        bp1.font.bold = True
        bp1.font.color.rgb = CHARCOAL
        
        bp2 = btf.add_paragraph()
        bp2.text = desc
        bp2.font.size = Pt(12)
        bp2.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 9: ARCHITECTURE & TECH STACK
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide9)
    add_card(slide9, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), WHITE)
    add_badge(slide9, Inches(1.3), Inches(0.95), "SYSTEM ARCHITECTURE", LAVENDER, CHARCOAL)

    h_box9 = slide9.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.5), Inches(0.8))
    hp9 = h_box9.text_frame.paragraphs[0]
    hp9.text = "Modern, Resilient & Scalable Tech Stack"
    hp9.font.size = Pt(30)
    hp9.font.bold = True
    hp9.font.color.rgb = CHARCOAL

    tech_blocks = [
        ("🎨 Frontend Experience", "React 19 & TypeScript\nVite Build Engine (<1s builds)\nTailwindCSS v4 Design System\nLucide Icons & Fluid Animations", YELLOW),
        ("⚙️ Backend Services", "FastAPI Asynchronous API\nPython 3.14 High-Speed Runtime\nDeterministic Scoring Engines\n13/13 Automated Tests Passing", LAVENDER),
        ("💾 Data & AI Engines", "SQLite Persistent Storage\nPyMuPDF Binary Parser\nTF-IDF Semantic Alias Matcher\nMulti-User Schema Isolation", SKY),
        ("🛡️ Security & Privacy", "PBKDF2 SHA-256 Salt Hashing\nProtected Route Auth Guards\nAnti-Bypass Verification Gates\n100% Student-First Privacy", YELLOW)
    ]
    tw = Inches(2.55)
    for idx, (title, details, col) in enumerate(tech_blocks):
        t_left = Inches(1.3) + idx * Inches(2.72)
        add_card(slide9, t_left, Inches(2.4), tw, Inches(3.9), col)
        
        box = slide9.shapes.add_textbox(t_left + Inches(0.15), Inches(2.55), tw - Inches(0.3), Inches(3.6))
        btf = box.text_frame
        btf.word_wrap = True
        
        p1 = btf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = CHARCOAL
        
        p2 = btf.add_paragraph()
        p2.text = f"\n{details}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 10: CONCLUSION & SUMMARY
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_canvas_background(slide10)
    add_card(slide10, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3), CHARCOAL)

    add_badge(slide10, Inches(1.3), Inches(1.1), "THE ROAD AHEAD", YELLOW, CHARCOAL)

    cbox = slide10.shapes.add_textbox(Inches(1.3), Inches(1.7), Inches(10.7), Inches(2.4))
    ctf = cbox.text_frame
    ctf.word_wrap = True
    
    cp1 = ctf.paragraphs[0]
    cp1.text = "Building Confidence, One Skill at a Time"
    cp1.font.size = Pt(44)
    cp1.font.bold = True
    cp1.font.color.rgb = WHITE
    
    cp2 = ctf.add_paragraph()
    cp2.text = "PathPilot replaces career confusion with transparent, joyful, and verifiable growth."
    cp2.font.size = Pt(18)
    cp2.font.color.rgb = AMBER

    # 3 Summary Highlights in White/Pastel Cards
    summary_highlights = [
        ("🌟 Humanized Design", "Ronas IT signature pastel cards, fluid transitions, and lovable Pip mascot.", YELLOW),
        ("🎯 Real-World Value", "Clear pathways to Junior AI/ML, Data Analyst, and Full-Stack roles.", LAVENDER),
        ("🚀 Live & Tested", "Fully operational web application running at http://127.0.0.1:5173.", SKY)
    ]
    sw_10 = Inches(3.45)
    for idx, (title, desc, col) in enumerate(summary_highlights):
        s_left = Inches(1.3) + idx * Inches(3.62)
        add_card(slide10, s_left, Inches(4.3), sw_10, Inches(1.9), col)
        
        box = slide10.shapes.add_textbox(s_left + Inches(0.15), Inches(4.45), sw_10 - Inches(0.3), Inches(1.6))
        btf = box.text_frame
        btf.word_wrap = True
        
        p1 = btf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = CHARCOAL
        
        p2 = btf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = CHARCOAL

    output_path = os.path.join(os.getcwd(), "PathPilot_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
